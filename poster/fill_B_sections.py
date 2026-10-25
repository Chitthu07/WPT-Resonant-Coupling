"""B-P1, B-P2: fill Member B's boxes in poster/EM_WPT_poster_template.pptx.

Run from the repo root, after `python poster/make_section5_figures.py`:
    python poster/fill_B_sections.py

Only boxes named "S2 ..." and "S5 ..." are changed:
  S2 body                       executive summary (text from poster/B-P2_executive_summary.md), 28 pt
  S5 fig_* / S5 field map       the slot rectangle is replaced by the PNG of the same size
  S5 fig_* caption              one-line takeaway, 24 pt
  S5 setup box                  design values, 24 pt
Formatting follows A-P1/A-P2: Calibri, near-black 1A1A1A, symbols in Cambria Math italic,
exact line spacing (28 pt text on 34 pt lines, 24 pt text on 29 pt lines).
Running it again replaces what it wrote before, so it is safe to rerun after a change.
"""
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "poster" / "EM_WPT_poster_template.pptx"
SUMMARY_MD = ROOT / "poster" / "B-P2_executive_summary.md"
BLACK, NAVY = RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x1B, 0x2A, 0x4A)

FIGURES = {                      # box name -> PNG made by make_section5_figures.py
    "S5 fig_k_vs_d": "fig_s5_k_vs_d.png",
    "S5 fig_eta_vs_d": "fig_s5_eta_vs_d.png",
    "S5 fig_eta_vs_f": "fig_s5_eta_vs_f.png",
    "S5 fig_eta_vs_RL": "fig_s5_eta_vs_RL.png",
    "S5 field map": "fig_s5_B_map.png",
}
# Mini-markup for captions: {x} = symbol in Cambria Math italic, _{x} = subscript, ^{x} = superscript.
CAPTIONS = {
    "S5 fig_k_vs_d caption":
        "Closed form and Neumann agree to ≤ 0.05 % (0.5–15 cm); {k} ∝ 1/{d}^{3} only for {d} ≫ {a}.",
    "S5 fig_eta_vs_d caption":
        "{η}_{max} is 97 % at 1 cm, 86 % at 5 cm, 57 % at 10 cm; the fixed 6 Ω load gives 44 % at 10 cm.",
    "S5 fig_eta_vs_f caption":
        "At 2 and 5 cm the power peak splits in two; at 10 cm a single sharp peak sits at {f}_{0}.",
    "S5 fig_eta_vs_RL caption":
        "The best load {R}_{L,opt} = {R}_{2}√(1+{x}) moves with spacing; 6 Ω is best only at 5 cm.",
}
SETUP = [  # (text, bold, colour)
    ("Setup", True, NAVY),
    ("Planar spiral, {N} = 10", False, BLACK),
    ("{d}_{out} = 100 mm", False, BLACK),
    ("wire 0.7 mm, pitch 1.0 mm", False, BLACK),
    ("{L} = 17.13 µH, {C} = 5.915 nF", False, BLACK),
    ("{f}_{0} = 500 kHz", False, BLACK),
    ("{R} = 0.455 Ω, {Q} = 118.3", False, BLACK),
    ("{R}_{L} = 6.0 Ω", False, BLACK),
    ("{V}_{s} = 5 V peak, {R}_{s} = 0", False, BLACK),
]
MAP_CAPTION = ("S5 field map caption", "|{B}| (µT) in the {x}–{z} plane: primary at 1 A, coil 2 at {d} = 3 cm.", 22)
SUMMARY_SYMBOLS = (("k", "M", "Q"))      # single letters set as math symbols in the summary


def tokens(text):
    """Split mini-markup into (text, kind) with kind in plain / math / sub / sup."""
    out, i = [], 0
    for m in re.finditer(r"(_|\^)?\{([^}]*)\}", text):
        if m.start() > i:
            out.append((text[i:m.start()], "plain"))
        kind = {"_": "sub", "^": "sup", None: "math"}[m.group(1)]
        out.append((m.group(2), kind))
        i = m.end()
    if i < len(text):
        out.append((text[i:], "plain"))
    return out


def add_run(par, text, size, colour, bold=False, kind="plain"):
    r = par.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic = Pt(size), bold, kind == "math"
    f.color.rgb = colour
    rpr = r._r.get_or_add_rPr()
    if kind == "math":
        f.name = "Cambria Math"
        for tag in ("a:ea", "a:cs"):
            el = rpr.makeelement(qn(tag), {"typeface": "Cambria Math"})
            rpr.append(el)
    else:
        f.name = "Calibri"
    if kind in ("sub", "sup"):
        rpr.set("baseline", "-25000" if kind == "sub" else "30000")
    return r


def set_spacing(par, line_pt, after_pt):
    par.line_spacing = Pt(line_pt)
    par.space_before, par.space_after = Pt(0), Pt(after_pt)


def fill_text(shape, lines, size, line_pt, after_pt):
    """lines: list of (markup text, bold, colour). Replaces all paragraphs of the box."""
    tf = shape.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for n, (text, bold, colour) in enumerate(lines):
        par = first if n == 0 else tf.add_paragraph()
        set_spacing(par, line_pt, after_pt)
        for t, kind in tokens(text):
            add_run(par, t, size, colour, bold=bold, kind=kind)


def summary_text():
    md = SUMMARY_MD.read_text(encoding="utf-8")
    body = md.split("\n\n", 2)[1].strip()          # the paragraph after the "# B-P2" heading
    for s in SUMMARY_SYMBOLS:                       # k, M, Q as math symbols (whole words only)
        body = re.sub(rf"(?<![\w{{]){s}(?![\w}}])", "{" + s + "}", body)
    return body.replace("f₀", "{f}_{0}")


def place_picture(slide, name, png):
    shape = next(s for s in slide.shapes if s.name == name)
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    shape._element.getparent().remove(shape._element)
    pic = slide.shapes.add_picture(str(ROOT / "poster" / png), left, top, width, height)
    pic.name = name


def main():
    prs = Presentation(PPTX)
    slide = prs.slides[0]
    names = {s.name for s in slide.shapes}
    missing = [n for n in [*FIGURES, *CAPTIONS, "S2 body", "S5 setup box"] if n not in names]
    if missing:
        sys.exit(f"boxes not found in the template: {missing}")
    by_name = {s.name: s for s in slide.shapes}

    fill_text(by_name["S2 body"], [(summary_text(), False, BLACK)], 28, 34, 8)
    for name, text in CAPTIONS.items():
        fill_text(by_name[name], [(text, False, BLACK)], 24, 29, 0)
    fill_text(by_name["S5 setup box"], SETUP, 24, 29, 0)
    name, text, size = MAP_CAPTION                            # one-line caption under the field map
    cap = by_name.get(name)
    if cap is None:
        m = by_name["S5 field map"]
        cap = slide.shapes.add_textbox(m.left, m.top + m.height + Inches(0.05), m.width, Inches(0.55))
        cap.name = name
        cap.text_frame.word_wrap = True
    fill_text(cap, [(text, False, BLACK)], size, 26, 0)
    for name, png in FIGURES.items():
        place_picture(slide, name, png)
    prs.save(PPTX)
    print(f"Filled S2 and S5 in {PPTX.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
