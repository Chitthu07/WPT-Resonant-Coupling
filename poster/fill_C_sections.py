"""C-P1: fill Member C's boxes (Sections 6 and 7) in poster/EM_WPT_poster_template.pptx.

Run from the repo root, after `python poster/make_c_p1_figure.py`:
    python poster/fill_C_sections.py
Text is poster/C-P1_final.md (numbers from poster/C-P1_numbers.md, R1 = R2 = 0.455 ohm).
Only shapes whose names start with "S6 " or "S7 " are touched; running it again replaces what it wrote before.
Formatting follows B's fill_B_sections.py: Calibri, near-black 1A1A1A, symbols in Cambria Math italic,
exact line spacing; mini-markup {x} = symbol, _{x} = subscript, ^{x} = superscript, **...** = bold.
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "poster"))
from fill_B_sections import tokens, add_run, set_spacing          # noqa: E402  (B's helpers, same style)

PPTX = ROOT / "poster" / "EM_WPT_poster_template.pptx"
FIG = ROOT / "figs" / "fig_c_p1_validation_chain.png"
BLACK, NAVY, WINE, WHITE = (RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x1B, 0x2A, 0x4A),
                            RGBColor(0x9B, 0x1D, 0x3A), RGBColor(0xFF, 0xFF, 0xFF))
LEFT, WIDTH = 31.84, 13.85                                        # S6/S7 content column [in]

S6_BULLETS = [
    "**No prototype was built:** closed form ↔ Python Neumann sum + 2 × 2 phasor solver ↔ LTspice sim/wpt.asc. "
    "LTspice agrees with the solver to 5 × 10^{−10} % (2, 5, 10 cm, 350–650 kHz).",
    "Faraday in time: at {f}_{0}, {V}_{2} = −j{ω}{M}{I}_{1} and {Z}_{2} = {R}_{2} + {R}_{L} is real, so {i}_{2} lags "
    "{i}_{1} by 90° ({T}_{0}/4 = 0.50 µs).",
    "Light-load peaks {f}_{±} (eq. 5) give {k} within 0.3 % at 1, 2, 5 cm (5 cm: 474.4 / 530.3 kHz → 0.1109).",
    "**Finding:** a 6 Ω load pulls them in (5 cm: 483.5 / 519.6 kHz → {k} = 0.072, 35 % low; 10 cm: 1 peak).",
]
S6_TABLE = [
    ["d (cm)", "M (µH)", "k", "ηmax (RL,opt)", "η at 6 Ω", "PL at 6 Ω (W)", "Regime"],
    ["1", "8.79", "0.513", "96.8 % (27.6 Ω)", "92.6 %", "0.098", "over-coupled"],
    ["3", "3.72", "0.217", "92.5 % (11.7 Ω)", "91.0 %", "0.526", "over-coupled"],
    ["5", "1.90", "0.111", "85.9 % (6.00 Ω)", "85.9 %", "1.79", "over-coupled; design point"],
    ["10", "0.51", "0.030", "57.2 % (1.67 Ω)", "43.5 %", "6.36", "near critical (k < kc)"],
]
S6_COLS = [1.0, 1.1, 0.9, 2.3, 1.3, 1.85, 5.4]                     # column widths [in], sum 13.85
S6_NOTES = [
    "{V}_{s} = 5 V pk, {R}_{s} = 0, {R}_{1} = {R}_{2} = 0.455 Ω; regime from {k}_{c} = 0.0318 ({R}_{L} = 6 Ω, {R}_{s} = 0).",
    "{k}, {M}: closed form vs Neumann agree to ≤ 0.05 % (0.5–15 cm).",
]
S6_CAPTION = ("(a) Model chain. (b) {P}_{L} (W) vs {f} (kHz) at 2, 5, 10 cm, {R}_{L} = 6 Ω; markers: LTspice; dotted lines: light-load "
              "{f}_{±} at 5 cm. (c) {i}_{1}, {i}_{2} (A) vs {t} (µs), 5 cm: {i}_{2} lags {i}_{1} by 90°. "
              "Read {k} only from light-load peaks.")
S7_BODY = [  # (markup, kind) kind: head or bullet
    ("Key findings", "head"),
    ("{Q} = 118.3: {η}_{max} = 96.8 % at 1 cm, 85.9 % at 5 cm ({R}_{L,opt} = 6.0 Ω), 57.2 % at 10 cm.", "bullet"),
    ("Max {η} ≠ max power: at 5 cm, {R}_{L,opt} = 6.0 Ω, but {R}_{L} = 79.1 Ω gives the most power.", "bullet"),
    ("Limitations", "head"),
    ("Range: {M} ∝ 1/{d}^{3} only for {d} ≫ {a} (log-log slope −2.1 at 8–10 cm, −2.97 at 50–100 cm).", "bullet"),
    ("AC loss: {R}_{dc} 0.13 → 0.455 Ω (skin 0.269, prox. 0.136, ESR 0.035, joints 0.015); range 0.38–0.48 Ω.", "bullet"),
    ("Misalignment (3 cm): {η} = 91.1 / 88.8 / 84.2 % at Δ{x} = 2 / 3 / 4 cm; {M} = 0 at 8.22 cm, then {M} < 0.", "bullet"),
    ("Al plate 5 mm behind coil 2: {k} −56.4 %, {R}_{2} +178 mΩ, {L}_{2} −51 %; ferrite: {k} +34.5 %. {η} (3 cm): 92.5 → 24.0 % (63.9 % with {C}_{2} retuned).", "bullet"),
    ("Stress ({R}_{s} = 0): {V}_{C1} 44.8 / 314 V pk at 5 / 10 cm; {I}_{1} ≤ 11.0 A pk, 27.5 W (50 Ω source: 0.1 A). "
     "Field: 82.3 µT rms at {z} = 0, < 27 µT rms (ICNIRP 2010) beyond 4.7 cm, < 1.84 beyond 15.4 cm.", "bullet"),
    ("**Future work:** build and measure a prototype; ferrite backing; adaptive frequency tuning [2].", "plain"),
]
S7_REFS = [
    "[1] A. Kurs et al., Science, vol. 317, no. 5834, pp. 83–86, 2007.",
    "[2] A. P. Sample, D. A. Meyer, and J. R. Smith, IEEE Trans. Ind. Electron., vol. 58, no. 2, pp. 544–554, 2011.",
    "[3] S. Y. R. Hui, W. Zhong, and C. K. Lee, IEEE Trans. Power Electron., vol. 29, no. 9, pp. 4500–4511, 2014.",
    "[4] R. K. Shevgaonkar, Electromagnetic Waves. New Delhi: Tata McGraw-Hill, 2006.",
]
# vertical layout [in] (top, height); S6 panel ends at 23.46 in, S7 panel at 32.42 in
LAYOUT = {"S6 bullets": (14.62, 2.25), "S6 table": (16.95, 2.05), "S6 notes": (19.05, 0.62),
          "S6 figure": (19.72, 2.75), "S6 caption": (22.52, 0.85),
          "S7 body": (25.15, 5.80), "S7 references": (31.00, 1.35)}


def write_par(par, markup, size, colour, bold=False):
    """Add runs for markup with **bold** spans and math/sub/sup tokens."""
    for n, chunk in enumerate(markup.split("**")):
        for t, kind in tokens(chunk):
            add_run(par, t, size, colour, bold=bold or n % 2 == 1, kind=kind)


def textbox(slide, name, lines, size, line_pt, after_pt, bullets=False, colour=BLACK):
    top, height = LAYOUT[name]
    tb = slide.shapes.add_textbox(Inches(LEFT), Inches(top), Inches(WIDTH), Inches(height))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for n, item in enumerate(lines):
        text, kind = item if isinstance(item, tuple) else (item, "bullet" if bullets else "plain")
        par = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        set_spacing(par, line_pt, after_pt)
        if kind == "head":
            write_par(par, text, size, WINE, bold=True)
        else:
            if kind == "bullet":
                pPr = par._p.get_or_add_pPr()
                pPr.set("marL", str(int(Inches(0.32))))
                pPr.set("indent", str(-int(Inches(0.32))))
                add_run(par, "•\t", size, colour)
            write_par(par, text, size, colour)
    return tb


def table(slide):
    top, height = LAYOUT["S6 table"]
    rows, cols = len(S6_TABLE), len(S6_TABLE[0])
    gt = slide.shapes.add_table(rows, cols, Inches(LEFT), Inches(top), Inches(WIDTH), Inches(height))
    gt.name = "S6 table"
    t = gt.table
    for j, w in enumerate(S6_COLS):
        t.columns[j].width = Inches(w)
    for i in range(rows):
        t.rows[i].height = Inches(height / rows)
        for j in range(cols):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(0xF2, 0xF4, 0xF8) if i % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            par = tf.paragraphs[0]
            set_spacing(par, 26, 0)
            add_run(par, S6_TABLE[i][j], 22, WHITE if i == 0 else BLACK, bold=i == 0)


def main():
    prs = Presentation(PPTX)
    slide = prs.slides[0]
    old = ("S6 no-prototype caption", "S6 comparison table", "S6 figure", *LAYOUT)
    for sh in list(slide.shapes):                       # remove the placeholders and anything written before
        if sh.name in old:
            sh._element.getparent().remove(sh._element)
    textbox(slide, "S6 bullets", S6_BULLETS, 24, 29, 3, bullets=True)
    table(slide)
    textbox(slide, "S6 notes", S6_NOTES, 20, 24, 0)
    top, height = LAYOUT["S6 figure"]
    pic = slide.shapes.add_picture(str(FIG), Inches(LEFT), Inches(top), width=Inches(WIDTH))
    pic.name = "S6 figure"
    textbox(slide, "S6 caption", [S6_CAPTION], 22, 26, 0)
    textbox(slide, "S7 body", S7_BODY, 24, 29, 3)
    textbox(slide, "S7 references", S7_REFS, 20, 24, 0)
    prs.save(PPTX)
    print(f"Saved {PPTX.relative_to(ROOT)}; figure height {pic.height / 914400:.2f} in")


if __name__ == "__main__":
    main()
