"""Put the team logo (media/team_logo_circle.png, circular crop with transparent corners) into the
"Logo (optional)" slot of poster/EM_WPT_poster_template.pptx, as a square centred in the slot.

Run from the repo root:  python poster/place_logo.py   (rerunnable: replaces the logo it placed before)
"""

from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "poster" / "EM_WPT_poster_template.pptx"
LOGO = ROOT / "media" / "team_logo_circle.png"
SLOT, PLACED = "Logo (optional)", "Logo"


def main():
    prs = Presentation(PPTX)
    slide = prs.slides[0]
    shapes = {s.name: s for s in slide.shapes}
    ref = shapes.get(SLOT) or shapes.get(PLACED)
    if ref is None:
        raise SystemExit(f"neither '{SLOT}' nor '{PLACED}' found in the template")
    left, top, width, height = ref.left, ref.top, ref.width, ref.height
    side = min(width, height)
    ref._element.getparent().remove(ref._element)
    pic = slide.shapes.add_picture(str(LOGO), left + (width - side) // 2, top + (height - side) // 2, side, side)
    pic.name = PLACED
    prs.save(PPTX)
    print(f"Logo placed: {side / 914400:.2f} x {side / 914400:.2f} in")


if __name__ == "__main__":
    main()
